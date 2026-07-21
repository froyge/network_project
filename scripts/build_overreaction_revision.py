from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Cm, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.utils import get_column_letter

OUT = Path('output')
OUT.mkdir(exist_ok=True)

NAVY = '17365D'
BLUE = '1F4E78'
LIGHT = 'D9EAF7'
PALE = 'EEF5FB'
GREEN = 'E2F0D9'
YELLOW = 'FFF2CC'
WHITE = 'FFFFFF'
GRAY = '666666'
THIN = Side(style='thin', color='B7C9D6')


def shade_cell(cell, color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), color)
    tc_pr.append(shd)


def set_cell_text(cell, text, bold=False, color=None, size=9):
    cell.text = ''
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(str(text))
    run.bold = bold
    run.font.name = 'Times New Roman'
    run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = __import__('docx').shared.RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_equation(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(text)
    r.font.name = 'Cambria Math'
    r._element.rPr.rFonts.set(qn('w:eastAsia'), 'Cambria Math')
    r.font.size = Pt(11)
    return p


def add_body(doc, text, bold_prefix=None):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.first_line_indent = Cm(1.25)
    p.paragraph_format.line_spacing = 1.25
    p.paragraph_format.space_after = Pt(4)
    if bold_prefix and text.startswith(bold_prefix):
        r1 = p.add_run(bold_prefix)
        r1.bold = True
        r2 = p.add_run(text[len(bold_prefix):])
        runs = [r1, r2]
    else:
        runs = [p.add_run(text)]
    for r in runs:
        r.font.name = 'Times New Roman'
        r._element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
        r.font.size = Pt(12)
    return p


def add_doc_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = 'Table Grid'
    for j, h in enumerate(headers):
        cell = table.rows[0].cells[j]
        shade_cell(cell, BLUE)
        set_cell_text(cell, h, bold=True, color=WHITE, size=8)
    for row in rows:
        cells = table.add_row().cells
        for j, value in enumerate(row):
            set_cell_text(cells[j], value, size=8)
            if j == 0:
                shade_cell(cells[j], PALE)
    if widths:
        for row in table.rows:
            for j, w in enumerate(widths):
                row.cells[j].width = Cm(w)
    doc.add_paragraph()
    return table


# ----------------------------- Word document -----------------------------
doc = Document()
section = doc.sections[0]
section.top_margin = Cm(2)
section.bottom_margin = Cm(2)
section.left_margin = Cm(2.5)
section.right_margin = Cm(1.5)

styles = doc.styles
styles['Normal'].font.name = 'Times New Roman'
styles['Normal']._element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
styles['Normal'].font.size = Pt(12)
for style_name in ['Title', 'Heading 1', 'Heading 2', 'Heading 3']:
    styles[style_name].font.name = 'Times New Roman'
    styles[style_name]._element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run('МОДЕЛЬ ПОВЕДЕНЧЕСКОЙ КАЛИБРОВКИ РЫНОЧНОЙ СТОИМОСТИ АКЦИЙ ПРИ ЭВРИСТИКЕ ИЗБЫТОЧНОЙ РЕАКЦИИ')
r.bold = True
r.font.name = 'Times New Roman'
r._element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
r.font.size = Pt(14)

add_body(doc, 'Настоящая версия модели исправляет три принципиальных недостатка предварительного варианта. Во-первых, эвристика избыточной реакции диагностируется строго по принятому в исследовании правилу: абсолютное изменение котировки в день события относительно предыдущего торгового дня составляет не менее 5%. Во-вторых, коэффициенты модели не задаются экспертно и не переносятся из другой выборки, а оцениваются только на обучающей части эмпирического массива. В-третьих, пригодность модели проверяется на отдельной тестовой части, которая не используется при определении коэффициентов.')

h = doc.add_heading('1. Что именно исправлено', level=1)
add_body(doc, 'Ранее в качестве зависимой переменной использовалось увеличение расстояния между котировкой и стоимостью по модели Гордона. Этот показатель имеет экономический смысл как дополнительная проверка направления ценового движения, но не совпадает с принятым в исследовании определением эвристики. Поэтому он исключается из основного диагностического критерия и может сохраняться только в проверках устойчивости.')
add_body(doc, 'Основной признак избыточной реакции теперь формируется независимо от модели стоимости. Сначала фиксируется факт экстремальной реакции рынка. Затем для уже идентифицированного наблюдения оценивается размер поведенческой составляющей отклонения котировки от фундаментального ориентира. Диагностика и стоимостная корректировка становятся двумя последовательными, но содержательно разными этапами.')

add_doc_table(doc,
    ['Элемент', 'Предварительный вариант', 'Исправленный вариант'],
    [
        ['Диагностика', 'Увеличение расстояния до стоимости по Гордону', '|P_t / P_(t-1) - 1| >= 5%'],
        ['Роль модели стоимости', 'Участвовала в определении факта эвристики', 'Используется только для измерения стоимостного отклонения'],
        ['Основная модель', 'Logit вероятности альтернативного критерия', 'Квантильная модель величины текущего отклонения'],
        ['Получение коэффициентов', 'Оценивание на всей выборке', 'Оценивание только на обучающей части'],
        ['Проверка', 'Внутривыборочные показатели', 'Хронологическая тестовая выборка и holdout по тикерам'],
        ['Результат', 'Вероятность и условная поправка', 'Нижняя, центральная и верхняя границы стоимости'],
    ], widths=[3.2, 6.0, 7.2])

h = doc.add_heading('2. Первый этап - диагностика эвристики', level=1)
add_body(doc, 'Для компании i в дату события t рассчитывается однодневная доходность акции:')
add_equation(doc, 'r_it = (P_it - P_i,t-1) / P_i,t-1.                                             (1)')
add_body(doc, 'Эвристика избыточной реакции диагностируется по бинарному критерию:')
add_equation(doc, 'OR_it = 1, если |r_it| >= 0,05; OR_it = 0, если |r_it| < 0,05.                   (2)')
add_body(doc, 'Критерий применяется непосредственно в день события. Последующее восстановление или разворот цены не требуется. Порог 5% является основным и полностью соответствует логике ранее выполненного исследования. Пороги 3%, 7,5% и 10% используются только в качестве проверок устойчивости, а не как альтернативное основное определение.')
add_body(doc, 'Для разделения положительной и отрицательной реакции вводятся переменные:')
add_equation(doc, 'OR_it^+ = 1{r_it >= 0,05};     OR_it^- = 1{r_it <= -0,05}.                       (3)')
add_body(doc, 'Интенсивность движения сверх установленной границы рассчитывается отдельно для роста и падения:')
add_equation(doc, 'I_it^+ = max(r_it - 0,05; 0) x 100;     I_it^- = max(-r_it - 0,05; 0) x 100.    (4)')
add_body(doc, 'Например, рост на 8% дает I+ = 3 процентных пункта, падение на 12% дает I- = 7 процентных пунктов.')

h = doc.add_heading('3. Второй этап - определение диапазона стоимости', level=1)
add_body(doc, 'Для каждого наблюдения рассчитывается фундаментальный стоимостной ориентир V_it^F. В эмпирической апробации используется модель Гордона, поскольку соответствующий показатель присутствует в исходном массиве. В прикладной оценке V_it^F может быть получена методом DCF, FCFE или путем согласования результатов нескольких подходов.')
add_equation(doc, 'Y_it = ln(P_it / V_it^F).                                                        (5)')
add_body(doc, 'Переменная Y показывает относительное отклонение текущей биржевой котировки от фундаментального ориентира. Положительное значение соответствует превышению котировки над расчетной стоимостью, отрицательное - нахождению котировки ниже нее.')
add_body(doc, 'Для формирования диапазона оцениваются три условных квантиля: 0,10, 0,50 и 0,90. Модель имеет следующий вид:')
add_equation(doc, 'Q_tau(Y_it | X_it) = a_tau + b_tau^+ OR_it^+ + b_tau^- OR_it^- + c_tau^+ I_it^+ + c_tau^- I_it^- + d_tau Beta_z,it + g_tau Event_it + s_tau Sector_i + l_tau Time_t.     (6)')

add_doc_table(doc,
    ['Переменная', 'Содержание', 'Функция'],
    [
        ['Y_it', 'ln(P_it / V_it^F)', 'Зависимая переменная - текущее стоимостное отклонение'],
        ['OR_it^+', 'Рост не менее 5%', 'Положительная ветвь эвристики'],
        ['OR_it^-', 'Падение не менее 5%', 'Отрицательная ветвь эвристики'],
        ['I_it^+', 'Превышение положительного движения над 5%', 'Интенсивность положительной реакции'],
        ['I_it^-', 'Превышение отрицательного движения над 5%', 'Интенсивность отрицательной реакции'],
        ['Beta_z,it', 'Стандартизированная beta', 'Контроль систематического риска'],
        ['Event_it', 'Дивидендная отсечка, МСФО, РСБУ, иная новость', 'Контроль событийного контекста'],
        ['Sector_i', 'Отраслевые фиктивные переменные', 'Контроль отраслевой неоднородности'],
        ['Time_t', 'Годовые эффекты или временной тренд', 'Контроль рыночного режима'],
    ], widths=[2.5, 7.0, 7.0])

h = doc.add_heading('4. Откуда берутся коэффициенты', level=1)
add_body(doc, 'Коэффициенты a, b, c, d, g, s и l не выбираются исследователем вручную. Для каждого квантиля они оцениваются по обучающей выборке методом квантильной регрессии. Оценка подбирает такие значения коэффициентов, при которых сумма асимметрично взвешенных абсолютных ошибок минимальна:')
add_equation(doc, 'b_hat_tau = arg min_b SUM_(i in Train) rho_tau(Y_i - X_i b),                     (7)')
add_equation(doc, 'rho_tau(u) = u[tau - 1{u < 0}].                                                  (8)')
add_body(doc, 'Коэффициенты для квантиля 0,50 описывают медианное отклонение. Коэффициенты для 0,10 и 0,90 задают нижнюю и верхнюю части условного распределения. В отличие от стандартного отклонения такой подход не требует нормальности распределения и допускает различную ширину диапазона в сторону повышения и понижения стоимости.')
add_body(doc, 'Logit-модель в основной расчет не включается. Причина состоит в том, что факт избыточной реакции уже известен из правила |r| >= 5%. Попытка прогнозировать этот же признак через величину доходности была бы круговой. Logit допустим только как вспомогательная модель ex ante, если вероятность экстремальной реакции оценивается до события по переменным, известным заранее. Для формирования текущего диапазона стоимости он не нужен.')

h = doc.add_heading('5. Разделение выборки и апробация', level=1)
add_body(doc, 'Выборка сортируется по дате. Первые 70% наблюдений образуют обучающую часть, последние 30% - тестовую. На обучающей части оцениваются все коэффициенты и правила объединения малочисленных отраслей. После этого коэффициенты фиксируются и без переоценки применяются к тестовым наблюдениям. Такая схема исключает использование будущей информации.')
add_body(doc, 'Дополнительно проводится проверка по новым эмитентам. Из выборки случайным образом, но на уровне тикеров, исключаются 20% компаний. Модель обучается на оставшихся эмитентах и проверяется на компаниях, данные которых не участвовали в определении коэффициентов.')

add_doc_table(doc,
    ['Блок', 'Назначение', 'Что рассчитывается'],
    [
        ['Обучающая выборка', 'Определение параметров модели', 'Коэффициенты Q10, Q50, Q90; правила укрупнения секторов'],
        ['Хронологический тест', 'Проверка на более позднем периоде', 'Pinball loss, MAE медианы, покрытие и ширина диапазона'],
        ['Holdout по тикерам', 'Проверка переноса на новые компании', 'Те же показатели на эмитентах вне обучения'],
        ['Секторальный тест', 'Проверка неоднородности', 'Показатели по секторам с достаточным числом наблюдений'],
        ['Событийный тест', 'Проверка вторичной роли событий', 'Результаты по дивидендам, МСФО, РСБУ и новостям'],
    ], widths=[3.5, 6.0, 7.0])

h = doc.add_heading('6. Получение диапазона рыночной стоимости', level=1)
add_body(doc, 'Поведенческая часть отклонения для конкретного наблюдения включает только коэффициенты, относящиеся к диагностированной эвристике и ее интенсивности:')
add_equation(doc, 'B_hat_tau,it = b_hat_tau^+ OR_it^+ + b_hat_tau^- OR_it^- + c_hat_tau^+ I_it^+ + c_hat_tau^- I_it^-.    (9)')
add_body(doc, 'Контрольные переменные не вычитаются из котировки. Они необходимы, чтобы коэффициенты поведенческого блока не поглощали отраслевые, временные и событийные особенности.')
add_body(doc, 'Текущая котировка очищается от оцененной поведенческой составляющей:')
add_equation(doc, 'V_hat_tau,it^corr = P_it x exp(-B_hat_tau,it).                                  (10)')
add_body(doc, 'Нижняя, центральная и верхняя границы определяются следующим образом:')
add_equation(doc, 'V_it^L = min(V_hat_0,10^corr; V_hat_0,90^corr);   V_it^C = V_hat_0,50^corr;   V_it^U = max(V_hat_0,10^corr; V_hat_0,90^corr).   (11)')
add_body(doc, 'Диапазон стоимости собственного капитала компании рассчитывается умножением на количество акций N_it:')
add_equation(doc, 'E_it^L = N_it V_it^L;     E_it^C = N_it V_it^C;     E_it^U = N_it V_it^U.        (12)')

h = doc.add_heading('7. Критерии качества', level=1)
add_body(doc, 'Для тестовой выборки оценивается не доля правильных классификаций, а качество стоимостного диапазона. Основными показателями являются медианная абсолютная ошибка центральной оценки, функция потерь квантильной регрессии, фактическое покрытие интервала Q10-Q90, средняя ширина диапазона и интервальный показатель Винклера. Для диапазона Q10-Q90 ожидаемое покрытие составляет около 80%. Избыточно широкая граница может обеспечивать высокое покрытие механически, поэтому покрытие рассматривается совместно со средней шириной.')

add_doc_table(doc,
    ['Показатель', 'Формула или принцип', 'Назначение'],
    [
        ['MAE медианы', 'median |Y - Q50|', 'Точность центральной оценки'],
        ['Pinball loss', 'rho_tau(Y - Q_tau)', 'Качество каждого квантиля'],
        ['Покрытие', 'Доля Y внутри [Q10; Q90]', 'Калибровка диапазона'],
        ['Средняя ширина', 'mean(Q90 - Q10)', 'Экономическая информативность диапазона'],
        ['Интервальный score', 'Ширина + штраф за непокрытие', 'Совместная оценка точности и компактности'],
        ['Benchmark', 'Секторное среднее +/- стандартное отклонение', 'Сопоставление с моделью статьи'],
    ], widths=[3.5, 6.0, 7.0])

h = doc.add_heading('8. Расширенная проверка устойчивости', level=1)
add_body(doc, 'Основной результат рассчитывается при пороге 5%. Для проверки устойчивости повторяются расчеты при порогах 3%, 7,5% и 10%; отдельно для роста и падения; по диапазонам силы реакции 5-7,5%, 7,5-10%, 10-15% и не менее 15%; по типам событий; по секторам; по временным режимам. В даты дивидендной отсечки дополнительно используется скорректированная доходность:')
add_equation(doc, 'r_it^adj = (P_it + D_it - P_i,t-1) / P_i,t-1.                                  (13)')
add_body(doc, 'Эта проверка отделяет механическое снижение котировки на величину дивиденда от поведенчески усиленной реакции.')

add_doc_table(doc,
    ['Срез', 'Основная спецификация', 'Проверка устойчивости'],
    [
        ['Порог', '5%', '3%; 7,5%; 10%'],
        ['Направление', 'Совместная модель с двумя ветвями', 'Рост и падение отдельно'],
        ['Интенсивность', 'Непрерывные I+ и I-', '5-7,5%; 7,5-10%; 10-15%; >=15%'],
        ['Событие', 'Контрольные переменные', 'Дивиденды, МСФО, РСБУ, новости отдельно'],
        ['Сектор', 'Фиксированные эффекты', 'Отдельные модели при достаточном N'],
        ['Время', 'Годовые эффекты', '2018-2021 и 2022-2025'],
        ['Дивидендная отсечка', 'Обычная доходность', 'Доходность с добавлением дивиденда'],
        ['Разбиение', 'Хронологические 70/30', 'Holdout 20% тикеров'],
        ['Интервал', 'Q10-Q90', 'Q05-Q95 и Q20-Q80'],
    ], widths=[3.0, 6.5, 7.0])

h = doc.add_heading('9. Ограничение текущего массива', level=1)
add_body(doc, 'Предоставленный массив содержит преимущественно наблюдения, уже отобранные по признаку абсолютного движения не менее 5%. На нем можно оценить условный диапазон стоимости внутри режима избыточной реакции и проверить перенос результатов на тестовую часть. Для оценки именно дополнительного эффекта OR относительно обычной реакции требуется расширить массив событиями, в которых |r| < 5%. В противном случае коэффициенты OR+ и OR- не могут быть отделены от общего уровня отклонения.')
add_body(doc, 'Поэтому в рабочей книге предусмотрены два режима. Режим A использует текущий массив и оценивает диапазон при условии, что OR уже диагностирована. Режим B является предпочтительным для диссертации и включает как экстремальные, так и обычные событийные наблюдения.')

h = doc.add_heading('10. Прикладная интерпретация', level=1)
add_body(doc, 'Модель используется в момент оценки. Сначала по изменению котировки устанавливается наличие эвристики. Затем по коэффициентам, заранее оцененным на обучающей выборке, определяется поведенческий диапазон. Если цена выросла не менее чем на 5%, положительная ветвь модели оценивает возможную поведенческую премию и корректирует текущую котировку вниз. Если цена снизилась не менее чем на 5%, отрицательная ветвь оценивает поведенческий дисконт и корректирует котировку вверх. Модель не утверждает, что цена впоследствии обязательно вернется к центральной оценке.')

h = doc.add_heading('11. Формулировка научной новизны', level=1)
add_body(doc, 'Разработана модель поведенческой калибровки рыночной стоимости акций при эвристике избыточной реакции, отличающаяся последовательным разделением моментной диагностики эффекта и количественной оценки его стоимостного последствия. Эвристика идентифицируется по превышению абсолютным изменением котировки порога 5% в день события, а величина поправки определяется по условным квантилям отклонения котировки от фундаментального стоимостного ориентира. Коэффициенты модели оцениваются на обучающей выборке и проверяются на независимой хронологической тестовой выборке и на эмитентах, не участвовавших в обучении. Результатом является нижняя, центральная и верхняя границы стоимости акции и собственного капитала компании, рассчитываемые непосредственно на дату проявления эвристики.')

h = doc.add_heading('12. Прикладные выводы', level=1)
add_body(doc, 'Первый вывод состоит в том, что факт избыточной реакции не требует вероятностной интерпретации: он наблюдается непосредственно через движение котировки не менее чем на 5%. Вероятностная или регрессионная модель нужна не для повторной диагностики этого факта, а для оценки величины стоимостного искажения.')
add_body(doc, 'Второй вывод связан с происхождением коэффициентов. Они являются результатом статистической оценки на исторической обучающей выборке. Тестовая выборка не участвует в их определении и используется только для проверки качества диапазона.')
add_body(doc, 'Третий вывод относится к роли событий. Дивидендная отсечка, МСФО и РСБУ включаются как контрольный контекст. Основной поведенческий блок образуют факт экстремальной реакции, ее знак и величина превышения порога.')
add_body(doc, 'Четвертый вывод имеет оценочное значение. Итогом модели является не прогноз будущей котировки и не механическое возвращение цены к модели Гордона, а диапазон текущей стоимости после исключения статистически оцененной поведенческой премии или дисконта.')

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
r = p.add_run('Источник формул и таблиц: разработано автором.')
r.italic = True
r.font.name = 'Times New Roman'
r.font.size = Pt(10)

doc_path = OUT / 'Модель_избыточной_реакции_исправленная.docx'
doc.save(doc_path)


# ----------------------------- Excel workbook -----------------------------
wb = Workbook()
readme = wb.active
readme.title = 'Описание'
input_ws = wb.create_sheet('Исходные данные')
split_ws = wb.create_sheet('Train_Test')
coef_ws = wb.create_sheet('Коэффициенты')
valid_ws = wb.create_sheet('Апробация')
slices_ws = wb.create_sheet('Срезы')
calc_ws = wb.create_sheet('Калькулятор')


def style_title(ws, title, end_col=10):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=end_col)
    c = ws.cell(1, 1, title)
    c.fill = PatternFill('solid', fgColor=NAVY)
    c.font = Font(name='Calibri', size=15, bold=True, color=WHITE)
    c.alignment = Alignment(vertical='center')
    ws.row_dimensions[1].height = 28


def style_header(row):
    for cell in row:
        cell.fill = PatternFill('solid', fgColor=BLUE)
        cell.font = Font(bold=True, color=WHITE)
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def style_range(ws, min_row, max_row, min_col, max_col):
    for row in ws.iter_rows(min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col):
        for c in row:
            c.border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
            c.alignment = Alignment(vertical='center', wrap_text=True)


# Description
style_title(readme, 'Исправленная модель поведенческой калибровки при избыточной реакции', 8)
readme['A3'] = 'Ключевое правило'
readme['B3'] = 'OR = 1, когда абсолютное изменение котировки в день события относительно предыдущего дня составляет не менее 5%.'
readme['A4'] = 'Основная модель'
readme['B4'] = 'Квантильная регрессия Q10, Q50, Q90 величины ln(P/VF).'
readme['A5'] = 'Обучение'
readme['B5'] = 'Первые 70% наблюдений после сортировки по дате.'
readme['A6'] = 'Тест'
readme['B6'] = 'Последние 30%; коэффициенты на тесте не переоцениваются.'
readme['A7'] = 'Результат'
readme['B7'] = 'Нижняя, центральная и верхняя границы стоимости акции и капитала.'
readme['A8'] = 'Важное ограничение'
readme['B8'] = 'Для оценки добавочного эффекта OR нужны и наблюдения с |r| < 5%. Если их нет, модель оценивает диапазон условно внутри режима OR.'
for row in range(3, 9):
    readme[f'A{row}'].fill = PatternFill('solid', fgColor=PALE)
    readme[f'A{row}'].font = Font(bold=True)
style_range(readme, 3, 8, 1, 2)
readme.column_dimensions['A'].width = 28
readme.column_dimensions['B'].width = 105

readme['A11'] = 'Используемые формулы'
readme['A11'].fill = PatternFill('solid', fgColor=BLUE)
readme['A11'].font = Font(bold=True, color=WHITE)
formulas = [
    ('Доходность', 'r = (P_t - P_t-1) / P_t-1'),
    ('Диагностика', 'OR = 1 при ABS(r) >= 5%'),
    ('Стоимостное отклонение', 'Y = LN(P_t / V_F)'),
    ('Положительная интенсивность', 'I+ = MAX(r - 5%; 0) x 100'),
    ('Отрицательная интенсивность', 'I- = MAX(-r - 5%; 0) x 100'),
    ('Скорректированная цена', 'V_corr,tau = P_t x EXP(-B_tau)'),
    ('Капитализация', 'E_tau = N x V_corr,tau'),
]
for idx, (name, form) in enumerate(formulas, start=12):
    readme.cell(idx, 1, name)
    readme.cell(idx, 2, form)
    readme.cell(idx, 1).fill = PatternFill('solid', fgColor=PALE)
    readme.cell(idx, 1).font = Font(bold=True)
style_range(readme, 12, 18, 1, 2)

# Input data template
headers = [
    'Ticker','Company','Date','Sector','Event_type','Price_t','Price_t_minus_1','Fundamental_value','Beta','Shares_mln',
    'Return','Abs_return','OR_5','OR_plus','OR_minus','Intensity_plus_pp','Intensity_minus_pp','Y_log_gap','Split','Reaction_band'
]
style_title(input_ws, 'Исходные данные и расчет признаков', len(headers))
for col, h in enumerate(headers, start=1):
    input_ws.cell(3, col, h)
style_header(input_ws[3])
for row in range(4, 204):
    input_ws.cell(row, 11, f'=IFERROR((F{row}-G{row})/G{row},"")')
    input_ws.cell(row, 12, f'=IF(K{row}="","",ABS(K{row}))')
    input_ws.cell(row, 13, f'=IF(L{row}="","",--(L{row}>=0.05))')
    input_ws.cell(row, 14, f'=IF(K{row}="","",--(K{row}>=0.05))')
    input_ws.cell(row, 15, f'=IF(K{row}="","",--(K{row}<=-0.05))')
    input_ws.cell(row, 16, f'=IF(K{row}="","",MAX(K{row}-0.05,0)*100)')
    input_ws.cell(row, 17, f'=IF(K{row}="","",MAX(-K{row}-0.05,0)*100)')
    input_ws.cell(row, 18, f'=IFERROR(LN(F{row}/H{row}),"")')
    input_ws.cell(row, 19, f'=IF(C{row}="","",IF(C{row}<=Train_Test!$B$5,"Train","Test"))')
    input_ws.cell(row, 20, f'=IF(L{row}="","",IF(L{row}<0.05,"<5%",IF(L{row}<0.075,"5-7.5%",IF(L{row}<0.1,"7.5-10%",IF(L{row}<0.15,"10-15%",">=15%")))))')
for col in range(1, len(headers)+1):
    input_ws.column_dimensions[get_column_letter(col)].width = 16
input_ws.column_dimensions['B'].width = 28
input_ws.column_dimensions['E'].width = 24
input_ws.freeze_panes = 'A4'
input_ws.auto_filter.ref = 'A3:T203'
style_range(input_ws, 3, 203, 1, len(headers))
for row in range(4, 204):
    input_ws.cell(row, 3).number_format = 'yyyy-mm-dd'
    input_ws.cell(row, 11).number_format = '0.00%'
    input_ws.cell(row, 12).number_format = '0.00%'
    input_ws.cell(row, 18).number_format = '0.0000'

# Split
style_title(split_ws, 'Разделение обучающей и тестовой выборок', 8)
split_ws['A3'] = 'Доля обучения'
split_ws['B3'] = 0.70
split_ws['B3'].number_format = '0%'
split_ws['A4'] = 'Минимальная дата'
split_ws['B4'] = '=MIN(\'Исходные данные\'!C4:C203)'
split_ws['A5'] = 'Дата границы 70%'
split_ws['B5'] = '=PERCENTILE.INC(\'Исходные данные\'!C4:C203,B3)'
split_ws['A6'] = 'Максимальная дата'
split_ws['B6'] = '=MAX(\'Исходные данные\'!C4:C203)'
split_ws['A8'] = 'N Train'
split_ws['B8'] = '=COUNTIF(\'Исходные данные\'!S4:S203,"Train")'
split_ws['A9'] = 'N Test'
split_ws['B9'] = '=COUNTIF(\'Исходные данные\'!S4:S203,"Test")'
split_ws['A11'] = 'Правило'
split_ws['B11'] = 'Коэффициенты Q10, Q50 и Q90 оцениваются только по строкам Train. На Test они применяются без переоценки.'
for row in [3,4,5,6,8,9,11]:
    split_ws[f'A{row}'].fill = PatternFill('solid', fgColor=PALE)
    split_ws[f'A{row}'].font = Font(bold=True)
style_range(split_ws, 3, 11, 1, 2)
split_ws.column_dimensions['A'].width = 28
split_ws.column_dimensions['B'].width = 90
split_ws['B4'].number_format = split_ws['B5'].number_format = split_ws['B6'].number_format = 'yyyy-mm-dd'

# Coefficients
style_title(coef_ws, 'Коэффициенты квантильной модели - заполняются результатами обучения', 8)
coef_headers = ['Переменная','Q10','Q50','Q90','SE / bootstrap','p-value','Экономический смысл','Источник']
for col, h in enumerate(coef_headers, start=1):
    coef_ws.cell(3, col, h)
style_header(coef_ws[3])
variables = [
    ('Intercept','','','','','','Базовый уровень отклонения','Оценка Train'),
    ('OR_plus','','','','','','Добавочный эффект роста >=5%','Оценка Train'),
    ('OR_minus','','','','','','Добавочный эффект падения <=-5%','Оценка Train'),
    ('Intensity_plus_pp','','','','','','Изменение эффекта на 1 п.п. сверх 5%','Оценка Train'),
    ('Intensity_minus_pp','','','','','','Изменение эффекта на 1 п.п. сверх 5%','Оценка Train'),
    ('Beta_z','','','','','','Контроль систематического риска','Оценка Train'),
    ('Event_dividend','','','','','','Контроль дивидендной отсечки','Оценка Train'),
    ('Event_IFRS','','','','','','Контроль МСФО','Оценка Train'),
    ('Event_RAS','','','','','','Контроль РСБУ','Оценка Train'),
    ('Event_news','','','','','','Контроль новостного контекста','Оценка Train'),
    ('Sector effects','','','','','','Отраслевые эффекты','Оценка Train'),
    ('Time effects','','','','','','Временной режим','Оценка Train'),
]
for r_idx, row in enumerate(variables, start=4):
    for c_idx, value in enumerate(row, start=1):
        coef_ws.cell(r_idx, c_idx, value)
    coef_ws.cell(r_idx, 1).fill = PatternFill('solid', fgColor=PALE)
    coef_ws.cell(r_idx, 1).font = Font(bold=True)
style_range(coef_ws, 3, 3+len(variables), 1, 8)
for col, width in enumerate([25,14,14,14,18,14,46,18], start=1):
    coef_ws.column_dimensions[get_column_letter(col)].width = width

coef_ws['A19'] = 'Как получаются коэффициенты'
coef_ws['A19'].fill = PatternFill('solid', fgColor=BLUE)
coef_ws['A19'].font = Font(bold=True, color=WHITE)
coef_ws.merge_cells('B19:H22')
coef_ws['B19'] = 'Для каждого квантиля коэффициенты минимизируют сумму pinball loss только на строках Train. Они не задаются вручную. Строки Test используются исключительно для проверки MAE, pinball loss, покрытия и ширины диапазона.'
coef_ws['B19'].alignment = Alignment(wrap_text=True, vertical='top')
style_range(coef_ws, 19, 22, 1, 8)

# Validation
style_title(valid_ws, 'Вневыборочная апробация', 9)
valid_headers = ['Модель / срез','N Train','N Test','MAE Q50','Pinball Q10','Pinball Q50','Pinball Q90','Покрытие Q10-Q90','Средняя ширина']
for col, h in enumerate(valid_headers, start=1):
    valid_ws.cell(3, col, h)
style_header(valid_ws[3])
validation_rows = [
    ['Основная модель, порог 5%','','','','','','','',''],
    ['Benchmark: секторное среднее +/- SD','','','','','','','',''],
    ['Holdout по тикерам','','','','','','','',''],
    ['Положительная реакция','','','','','','','',''],
    ['Отрицательная реакция','','','','','','','',''],
    ['Дивидендная отсечка','','','','','','','',''],
    ['МСФО','','','','','','','',''],
    ['РСБУ','','','','','','','',''],
    ['Новости','','','','','','','',''],
]
for r_idx, row in enumerate(validation_rows, start=4):
    for c_idx, value in enumerate(row, start=1):
        valid_ws.cell(r_idx, c_idx, value)
    valid_ws.cell(r_idx, 1).fill = PatternFill('solid', fgColor=PALE)
style_range(valid_ws, 3, 3+len(validation_rows), 1, 9)
for col in range(1, 10):
    valid_ws.column_dimensions[get_column_letter(col)].width = 18
valid_ws.column_dimensions['A'].width = 38
valid_ws['A15'] = 'Целевое покрытие интервала Q10-Q90'
valid_ws['B15'] = 0.80
valid_ws['B15'].number_format = '0%'
valid_ws['A16'] = 'Интерпретация'
valid_ws['B16'] = 'Покрытие оценивается совместно со средней шириной. Более широкий интервал не считается автоматически лучшим.'
valid_ws['A15'].fill = valid_ws['A16'].fill = PatternFill('solid', fgColor=YELLOW)
style_range(valid_ws, 15, 16, 1, 2)

# Slices
style_title(slices_ws, 'Срезы и проверки устойчивости', 8)
slice_headers = ['Группа проверки','Срез','Порог / условие','N Train','N Test','Q10','Q50','Q90']
for col, h in enumerate(slice_headers, start=1):
    slices_ws.cell(3, col, h)
style_header(slices_ws[3])
slice_rows = []
for threshold in ['3%','5% - основная','7.5%','10%']:
    slice_rows.append(['Порог',threshold,threshold,'','','','',''])
for direction in ['Рост','Падение']:
    slice_rows.append(['Направление',direction,direction,'','','','',''])
for band in ['5-7.5%','7.5-10%','10-15%','>=15%']:
    slice_rows.append(['Интенсивность',band,band,'','','','',''])
for event in ['Дивидендная отсечка','МСФО','РСБУ','Новость','Без события']:
    slice_rows.append(['Тип события',event,event,'','','','',''])
for period in ['2018-2021','2022-2025']:
    slice_rows.append(['Период',period,period,'','','','',''])
for method in ['Обычная доходность','С поправкой на дивиденд']:
    slice_rows.append(['Дивидендный тест',method,method,'','','','',''])
for r_idx, row in enumerate(slice_rows, start=4):
    for c_idx, value in enumerate(row, start=1):
        slices_ws.cell(r_idx, c_idx, value)
    slices_ws.cell(r_idx, 1).fill = PatternFill('solid', fgColor=PALE)
style_range(slices_ws, 3, 3+len(slice_rows), 1, 8)
for col in range(1, 9):
    slices_ws.column_dimensions[get_column_letter(col)].width = 18
slices_ws.column_dimensions['A'].width = 24
slices_ws.column_dimensions['B'].width = 28

# Calculator
style_title(calc_ws, 'Калькулятор текущей поведенческой корректировки', 8)
calc_ws['A3'] = 'Входные данные'
calc_ws['A3'].fill = PatternFill('solid', fgColor=BLUE)
calc_ws['A3'].font = Font(bold=True, color=WHITE)
calc_inputs = [
    ('Текущая цена акции',100.0),
    ('Цена предыдущего дня',94.0),
    ('Фундаментальная стоимость',90.0),
    ('Количество акций, млн',1000.0),
    ('Beta_z',0.0),
    ('Тип события','Новость'),
    ('Сектор',''),
    ('Год',2025),
]
for idx, (label, value) in enumerate(calc_inputs, start=4):
    calc_ws.cell(idx,1,label)
    calc_ws.cell(idx,2,value)
    calc_ws.cell(idx,1).fill = PatternFill('solid', fgColor=PALE)
    calc_ws.cell(idx,1).font = Font(bold=True)
    calc_ws.cell(idx,2).font = Font(color='0000FF')
style_range(calc_ws, 4, 11, 1, 2)

calc_ws['D3'] = 'Диагностика'
calc_ws['D3'].fill = PatternFill('solid', fgColor=BLUE)
calc_ws['D3'].font = Font(bold=True, color=WHITE)
calc_rows = [
    ('Доходность','=IFERROR((B4-B5)/B5,"")'),
    ('Абсолютная доходность','=ABS(E4)'),
    ('OR при 5%','=--(E5>=0.05)'),
    ('OR+','=--(E4>=0.05)'),
    ('OR-','=--(E4<=-0.05)'),
    ('I+ п.п.','=MAX(E4-0.05,0)*100'),
    ('I- п.п.','=MAX(-E4-0.05,0)*100'),
    ('Решение','=IF(E6=1,"Эвристика диагностирована","Эвристика не диагностирована")'),
]
for idx, (label, formula) in enumerate(calc_rows, start=4):
    calc_ws.cell(idx,4,label)
    calc_ws.cell(idx,5,formula)
    calc_ws.cell(idx,4).fill = PatternFill('solid', fgColor=PALE)
    calc_ws.cell(idx,4).font = Font(bold=True)
style_range(calc_ws, 4, 11, 4, 5)
calc_ws['E4'].number_format = calc_ws['E5'].number_format = '0.00%'

calc_ws['A14'] = 'Коэффициенты поведенческого блока'
calc_ws['A14'].fill = PatternFill('solid', fgColor=BLUE)
calc_ws['A14'].font = Font(bold=True, color=WHITE)
calc_ws['A15'] = 'Коэффициент'
calc_ws['B15'] = 'Q10'
calc_ws['C15'] = 'Q50'
calc_ws['D15'] = 'Q90'
style_header(calc_ws[15][:4])
for idx, var in enumerate(['OR_plus','OR_minus','Intensity_plus_pp','Intensity_minus_pp'], start=16):
    calc_ws.cell(idx,1,var)
    calc_ws.cell(idx,1).fill = PatternFill('solid', fgColor=PALE)
    # Link to coefficient sheet rows 5-8 based on current layout
    source_row = {'OR_plus':5,'OR_minus':6,'Intensity_plus_pp':7,'Intensity_minus_pp':8}[var]
    calc_ws.cell(idx,2,f'=Коэффициенты!B{source_row}')
    calc_ws.cell(idx,3,f'=Коэффициенты!C{source_row}')
    calc_ws.cell(idx,4,f'=Коэффициенты!D{source_row}')
style_range(calc_ws, 15, 19, 1, 4)

calc_ws['F14'] = 'Результат'
calc_ws['F14'].fill = PatternFill('solid', fgColor=BLUE)
calc_ws['F14'].font = Font(bold=True, color=WHITE)
result_rows = [
    ('B_Q10','=B16*E7+B17*E8+B18*E9+B19*E10'),
    ('B_Q50','=C16*E7+C17*E8+C18*E9+C19*E10'),
    ('B_Q90','=D16*E7+D17*E8+D18*E9+D19*E10'),
    ('Цена Q10','=IF(E6=1,$B$4*EXP(-G15),$B$4)'),
    ('Цена Q50','=IF(E6=1,$B$4*EXP(-G16),$B$4)'),
    ('Цена Q90','=IF(E6=1,$B$4*EXP(-G17),$B$4)'),
    ('Нижняя цена','=MIN(G18,G20)'),
    ('Центральная цена','=G19'),
    ('Верхняя цена','=MAX(G18,G20)'),
    ('Нижняя капитализация, млрд','=G21*$B$7/1000'),
    ('Центральная капитализация, млрд','=G22*$B$7/1000'),
    ('Верхняя капитализация, млрд','=G23*$B$7/1000'),
]
for idx, (label, formula) in enumerate(result_rows, start=15):
    calc_ws.cell(idx,6,label)
    calc_ws.cell(idx,7,formula)
    calc_ws.cell(idx,6).fill = PatternFill('solid', fgColor=PALE)
    calc_ws.cell(idx,6).font = Font(bold=True)
style_range(calc_ws, 15, 26, 6, 7)
calc_ws['F28'] = 'Примечание'
calc_ws['G28'] = 'Калькулятор становится расчетным после заполнения коэффициентов Q10, Q50 и Q90, оцененных только на Train.'
calc_ws['F28'].fill = PatternFill('solid', fgColor=YELLOW)
calc_ws['G28'].fill = PatternFill('solid', fgColor=YELLOW)
style_range(calc_ws, 28, 28, 6, 7)

for ws in wb.worksheets:
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = ws.freeze_panes or 'A3'
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is not None:
                cell.alignment = Alignment(vertical='center', wrap_text=True)
                if cell.row != 1 and cell.font == Font():
                    cell.font = Font(name='Calibri', size=10)

calc_ws.column_dimensions['A'].width = 32
calc_ws.column_dimensions['B'].width = 18
calc_ws.column_dimensions['C'].width = 15
calc_ws.column_dimensions['D'].width = 28
calc_ws.column_dimensions['E'].width = 24
calc_ws.column_dimensions['F'].width = 34
calc_ws.column_dimensions['G'].width = 28

xlsx_path = OUT / 'Модель_избыточной_реакции_исправленная.xlsx'
wb.save(xlsx_path)

print(doc_path)
print(xlsx_path)
